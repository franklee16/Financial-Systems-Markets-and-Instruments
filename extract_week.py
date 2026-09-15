"""Extract slide content and images from any week PPTX for Beamer conversion.

Usage: python extract_week.py <path-to-pptx>
Run from the target week subfolder (e.g. beamer/week2_InterestRate/).
Writes weekN_extract.md (structured slide dump) and figures/*.png.

Image recovery is relationship-based, not shape-based: any media file a slide
references is extracted, including image-fill placeholders that python-pptx's
shape walk misses (bit the week1 conversion on its debt/equity chart slides).
"""
import os
import re
import sys
import zipfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC = sys.argv[1]
week = os.path.basename(os.getcwd()).split("_")[0]
OUT_MD = f"{week}_extract.md"
FIG_DIR = "figures"


def extract_images_via_rels(z, slide_no, lines):
    """Pull every media target referenced by slideN.xml.rels."""
    rels_path = f"ppt/slides/_rels/slide{slide_no}.xml.rels"
    try:
        rels = z.read(rels_path).decode("utf-8")
    except KeyError:
        return 0
    media = re.findall(r'Target="\.\./media/([^"]+)"', rels)
    n = 0
    for m in media:
        ext = m.split(".")[-1].replace("jpeg", "jpg")
        out = os.path.join(FIG_DIR, f"slide{slide_no:02d}_img.{ext}")
        # multiple media on one slide get suffixed
        if os.path.exists(out):
            out = os.path.join(FIG_DIR, f"slide{slide_no:02d}_img{n}.{ext}")
        with open(out, "wb") as f:
            f.write(z.read(f"ppt/media/{m}"))
        lines.append(f"  [IMAGE -> {out.replace(os.sep, '/')}]")
        n += 1
    return n


def dump_text_frame(tf, lines):
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs).strip()
        if not text:
            continue
        indent = "  " * (p.level or 0)
        bullet = "-" if (p.level or 0) > 0 else ""
        lines.append(f"{indent}{bullet} {text}".rstrip())


def dump_shape(sh, lines):
    if sh.has_table:
        lines.append("  [TABLE]")
        for row in sh.table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            lines.append("    | " + " | ".join(cells) + " |")
        return
    if getattr(sh, "has_chart", False):
        lines.append("  [NATIVE CHART - must recreate or screenshot]")
        return
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in sh.shapes:
            dump_shape(sub, lines)
        return
    if sh.has_text_frame:
        dump_text_frame(sh.text_frame, lines)


def main():
    os.makedirs(FIG_DIR, exist_ok=True)
    prs = Presentation(SRC)
    z = zipfile.ZipFile(SRC)
    lines = [f"# Extraction: {SRC} ({len(prs.slides)} slides)", ""]
    total_imgs = 0
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"## Slide {i}  [layout: {slide.slide_layout.name}]")
        for sh in slide.shapes:
            dump_shape(sh, lines)
        total_imgs += extract_images_via_rels(z, i, lines)
        if slide.has_notes_slide:
            ntf = slide.notes_slide.notes_text_frame
            notes = ntf.text.strip() if ntf is not None else ""
            if notes:
                lines.append(f"  [NOTES: {notes[:400]}]")
        lines.append("")
    with open(OUT_MD, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Wrote {OUT_MD} ({len(prs.slides)} slides, {total_imgs} images)")


if __name__ == "__main__":
    main()
